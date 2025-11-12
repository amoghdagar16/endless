from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Iterable, Optional

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.exceptions import ValidationError
from app.db.models.account import Account
from app.db.models.types import AccountType
from app.services.audit import log_action


EXPECTED_COLUMNS = {
    "number",
    "name",
    "type",
    "detail_type",
    "parent_number",
    "is_active",
    "opening_balance",
    "opening_balance_date",
}

ACCOUNT_TYPES = {t.value for t in AccountType}
BALANCE_ALLOWED = {AccountType.ASSET.value, AccountType.LIABILITY.value, AccountType.EQUITY.value}


@dataclass
class ImportMessage:
    level: str
    message: str
    row: int | None = None


@dataclass
class ProposedAccount:
    number: str
    name: str
    type: str
    detail_type: Optional[str]
    parent_number: Optional[str]
    is_active: bool
    opening_balance: Optional[Decimal]
    opening_balance_date: Optional[date]
    source_row: int


@dataclass
class ImportResult:
    proposed_accounts: list[ProposedAccount]
    errors: list[ImportMessage]
    warnings: list[ImportMessage]
    inserted: int = 0


def list_accounts(session: Session, company_id: uuid.UUID) -> list[Account]:
    return (
        session.execute(
            select(Account).where(Account.company_id == company_id).order_by(Account.number)
        )
        .scalars()
        .all()
    )


def list_accounts_hierarchy(session: Session, company_id: uuid.UUID) -> list[dict[str, object]]:
    accounts = list_accounts(session, company_id)
    by_parent: dict[uuid.UUID | None, list[Account]] = {}
    for account in accounts:
        by_parent.setdefault(account.parent_id, []).append(account)

    def _build(parent_id: uuid.UUID | None) -> list[dict[str, object]]:
        nodes = []
        for acct in sorted(by_parent.get(parent_id, []), key=lambda a: a.number):
            nodes.append(
                {
                    "id": str(acct.id),
                    "number": acct.number,
                    "name": acct.name,
                    "type": acct.type.value if isinstance(acct.type, AccountType) else acct.type,
                    "detail_type": acct.detail_type,
                    "is_active": acct.is_active,
                    "children": _build(acct.id),
                }
            )
        return nodes

    return _build(None)


def export_accounts_csv(session: Session, company_id: uuid.UUID) -> str:
    accounts = list_accounts(session, company_id)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "number",
            "name",
            "type",
            "detail_type",
            "parent_number",
            "is_active",
            "opening_balance",
            "opening_balance_date",
        ]
    )
    id_to_number = {acct.id: acct.number for acct in accounts}

    for account in accounts:
        parent_number = id_to_number.get(account.parent_id) if account.parent_id else ""
        writer.writerow(
            [
                account.number,
                account.name,
                account.type.value if isinstance(account.type, AccountType) else account.type,
                account.detail_type or "",
                parent_number,
                str(account.is_active).lower(),
                account.opening_balance if account.opening_balance is not None else "",
                account.opening_balance_date.isoformat() if account.opening_balance_date else "",
            ]
        )
    return output.getvalue()


def parse_coa_bytes(data: bytes, filename: str) -> list[ProposedAccount]:
    if filename.lower().endswith(".csv"):
        return _parse_csv(data)
    if filename.lower().endswith(".xlsx"):
        return _parse_xlsx(data)
    raise ValidationError("Unsupported file type. Upload CSV or XLSX.")


def _parse_csv(data: bytes) -> list[ProposedAccount]:
    stream = io.StringIO(data.decode("utf-8-sig"))
    reader = csv.DictReader(stream)
    _ensure_columns(reader.fieldnames)
    accounts = []
    for idx, row in enumerate(reader, start=2):
        accounts.append(_normalize_row(row, idx))
    return accounts


def _parse_xlsx(data: bytes) -> list[ProposedAccount]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValidationError("Workbook is empty")
    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    _ensure_columns(headers)
    accounts: list[ProposedAccount] = []
    for idx, row in enumerate(rows[1:], start=2):
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        accounts.append(_normalize_row(row_dict, idx))
    return accounts


def _ensure_columns(headers: Iterable[str] | None) -> None:
    if not headers:
        raise ValidationError("File is missing headers")
    normalized = {str(h).strip().lower() for h in headers if h}
    missing = EXPECTED_COLUMNS - normalized
    if missing:
        raise ValidationError(f"Missing required columns: {', '.join(sorted(missing))}")


def _normalize_row(row: dict[str, object], row_index: int) -> ProposedAccount:
    def _get(name: str) -> str:
        value = row.get(name)
        if value is None:
            return ""
        if isinstance(value, str):
            return _collapse_whitespace(value)
        return str(value).strip()

    number = _get("number")
    name = _get("name")
    account_type = _get("type").lower()
    detail_type = _get("detail_type") or None
    parent_number = _get("parent_number") or None
    is_active = _to_bool(_get("is_active"))
    opening_balance = _to_decimal(_get("opening_balance"))
    opening_balance_date = _to_date(row.get("opening_balance_date"))

    return ProposedAccount(
        number=number,
        name=name,
        type=account_type,
        detail_type=detail_type,
        parent_number=parent_number,
        is_active=is_active,
        opening_balance=opening_balance,
        opening_balance_date=opening_balance_date,
        source_row=row_index,
    )


def validate_accounts(
    session: Session,
    *,
    company_id: uuid.UUID,
    accounts: Iterable[ProposedAccount],
) -> ImportResult:
    proposed = list(accounts)
    errors: list[ImportMessage] = []
    warnings: list[ImportMessage] = []

    existing_numbers = set(
        session.execute(select(Account.number).where(Account.company_id == company_id)).scalars()
    )

    seen: dict[str, ProposedAccount] = {}
    for account in proposed:
        if not account.number:
            errors.append(ImportMessage("error", "Account number is required", account.source_row))
        elif account.number in seen:
            errors.append(
                ImportMessage(
                    "error",
                    f"Duplicate account number {account.number}",
                    account.source_row,
                )
            )
        else:
            seen[account.number] = account

        if account.number in existing_numbers:
            errors.append(
                ImportMessage(
                    "error",
                    f"Account {account.number} already exists",
                    account.source_row,
                )
            )

        if account.type not in ACCOUNT_TYPES:
            errors.append(
                ImportMessage("error", f"Invalid account type '{account.type}'", account.source_row)
            )

        if account.opening_balance is not None:
            if account.type not in BALANCE_ALLOWED:
                errors.append(
                    ImportMessage(
                        "error",
                        "Opening balance allowed only for asset, liability, equity",
                        account.source_row,
                    )
                )
            if not account.opening_balance_date:
                errors.append(
                    ImportMessage(
                        "error",
                        "Opening balance date required when opening balance provided",
                        account.source_row,
                    )
                )

        # parent validation happens after all rows processed

    for account in proposed:
        if account.parent_number and account.parent_number in seen:
            parent = seen[account.parent_number]
            if parent.type != account.type:
                errors.append(
                    ImportMessage(
                        "error",
                        f"Parent type mismatch for {account.number} -> {account.parent_number}",
                        account.source_row,
                    )
                )

    if errors:
        return ImportResult(proposed_accounts=proposed, errors=errors, warnings=warnings)

    missing_parents = {
        acct.parent_number
        for acct in proposed
        if acct.parent_number and acct.parent_number not in seen and acct.parent_number not in existing_numbers
    }
    for parent in sorted(missing_parents):
        errors.append(ImportMessage("error", f"Parent account {parent} not found", None))

    if missing_parents:
        return ImportResult(proposed_accounts=proposed, errors=errors, warnings=warnings)

    return ImportResult(proposed_accounts=proposed, errors=[], warnings=warnings)


def apply_import(
    session: Session,
    *,
    company_id: uuid.UUID,
    accounts: Iterable[ProposedAccount],
    actor_id: Optional[uuid.UUID],
) -> ImportResult:
    validation = validate_accounts(session, company_id=company_id, accounts=accounts)
    if validation.errors:
        raise ValidationError("COA import failed", details=[msg.__dict__ for msg in validation.errors])

    ordered = _order_accounts(validation.proposed_accounts)

    existing_accounts = dict(
        session.execute(
            select(Account.number, Account.id).where(Account.company_id == company_id)
        ).all()
    )

    number_to_id: dict[str, uuid.UUID] = {}
    inserted = 0

    for account in ordered:
        if isinstance(account.type, str):
            normalized_type = account.type.strip().lower()
        elif isinstance(account.type, AccountType):
            normalized_type = account.type.value
        else:
            raise ValidationError(f"Unsupported account type '{account.type}'")

        try:
            account_type_enum = AccountType(normalized_type)
        except ValueError as exc:
            raise ValidationError(f"Invalid account type '{account.type}'") from exc

        db_account = Account(
            company_id=company_id,
            number=account.number,
            name=account.name,
            type=account_type_enum,
            detail_type=account.detail_type,
            is_active=account.is_active,
            opening_balance=account.opening_balance,
            opening_balance_date=account.opening_balance_date,
        )
        if account.parent_number:
            parent_id = number_to_id.get(account.parent_number) or existing_accounts.get(
                account.parent_number
            )
            if parent_id:
                db_account.parent_id = parent_id
        session.add(db_account)
        session.flush()
        number_to_id[account.number] = db_account.id
        inserted += 1

    log_action(
        session,
        company_id=company_id,
        actor_id=actor_id,
        action="accounts_imported",
        entity="account",
        entity_id=list(number_to_id.values())[-1] if number_to_id else uuid.uuid4(),
        before=None,
        after={"inserted": inserted},
    )

    session.commit()

    validation.inserted = inserted
    return validation


def _order_accounts(accounts: list[ProposedAccount]) -> list[ProposedAccount]:
    ordered: list[ProposedAccount] = []
    remaining = {acct.number: acct for acct in accounts}

    while remaining:
        progress = False
        for number, account in list(remaining.items()):
            if not account.parent_number or account.parent_number in {acct.number for acct in ordered}:
                ordered.append(account)
                remaining.pop(number)
                progress = True
        if not progress:
            raise ValidationError("Cyclic or missing parent relationships detected")
    return ordered


def _collapse_whitespace(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip())


def _to_bool(value: str) -> bool:
    if not value:
        return True
    return value.strip().lower() in {"true", "1", "yes", "y"}


def _to_decimal(value: str) -> Optional[Decimal]:
    if value is None or value == "" or (isinstance(value, str) and value.strip() == ""):
        return None
    try:
        cleaned = value.replace(",", "") if isinstance(value, str) else str(value)
        cleaned = cleaned.strip()
        if not cleaned:
            return None
        return Decimal(cleaned)
    except Exception as exc:  # pragma: no cover - defensive
        raise ValidationError(f"Invalid decimal value '{value}'") from exc


def _to_date(value: object) -> Optional[date]:
    if value in (None, ""):
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        value_str = value.strip()
        if not value_str:
            return None

        # Try ISO format first (YYYY-MM-DD)
        try:
            return date.fromisoformat(value_str)
        except ValueError:
            pass

        # Try common US formats: M/D/YY, M/D/YYYY, MM/DD/YY, MM/DD/YYYY
        for fmt in ['%m/%d/%y', '%m/%d/%Y', '%-m/%-d/%y', '%-m/%-d/%Y']:
            try:
                return datetime.strptime(value_str, fmt).date()
            except ValueError:
                continue

        raise ValidationError(f"Invalid date value '{value_str}'. Use YYYY-MM-DD or M/D/YY format")
    raise ValidationError(f"Unsupported date value: {value}")
